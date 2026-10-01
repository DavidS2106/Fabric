"""
reporters/preflight_package_reporter.py
───────────────────────────────────────
The preflight reports. Preflight has its own logical phase, so it gets its own
resolved run folder and never writes a Smoke report of any kind::

    output/Preflight/<run_id>_<timestamp>/{Excel,JSON,HTML}/...

Only the row shaping and the three renderings live here. The folder, the run
folder and the file names come from `report_output` through
:class:`~reporters.report_paths.ReportPathResolver`, exactly as they do for
every validation phase.
"""

from __future__ import annotations

import json
import logging
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from .report_paths import (
    ReportPathError,
    ResolvedReportPaths,
    filesystem_path,
)
from .preflight_html_template import render_preflight_html


logger = logging.getLogger(__name__)


#: Selector this module uses to shape preflight rows. It is the writer's own
#: vocabulary, not an output folder name: where the reports land is decided by
#: `report_output.phase_reports.preflight`.
PREFLIGHT_VALIDATION_TYPE = "preflight"

#: Heading used when the `reporting` block does not configure one.
DEFAULT_PREFLIGHT_TITLE = "Preflight Connection Validation"

PREFLIGHT_COLUMNS = (
    "run_id", "layer_name", "configuration_file_name", "table_configuration_id",
    "source_connection_reference", "target_connection_reference",
    "source_connection_type", "target_connection_type", "source_server_name",
    "target_server_name", "source_database_name", "target_database_name",
    "source_schema_name", "target_schema_name", "source_table_name",
    "target_table_name", "source_connection_validation_status",
    "target_connection_validation_status", "source_database_availability_status",
    "target_database_availability_status", "source_schema_availability_status",
    "target_schema_availability_status", "source_table_availability_status",
    "target_table_availability_status", "source_column_availability_status",
    "target_column_availability_status", "source_key_column_availability_status",
    "target_key_column_availability_status", "source_object_validation_status",
    "target_object_validation_status", "overall_validation_status",
    "failed_validation_checks", "validation_status", "message", "failure_reason",
    "validation_timestamp", "validation_start_time",
    "validation_completion_time", "validation_duration_seconds",
)

class SeparateValidationPackageWriter:
    """Render the preflight reports into an already resolved run folder.

    The class never decides where a report goes. It is handed the resolved
    :class:`~reporters.report_paths.ResolvedReportPaths` of the preflight
    phase, writes only the formats that layout enables, and reports which
    formats it wrote.
    """

    def __init__(self, reporting_settings: Optional[Dict[str, Any]] = None):
        self.reporting_settings = reporting_settings or {}

    @property
    def title(self) -> str:
        return str(
            self.reporting_settings.get("preflight_title")
            or DEFAULT_PREFLIGHT_TITLE
        )

    def target_directories(self, paths: ResolvedReportPaths) -> List[Path]:
        """Every directory this writer may create, for the shared-root check."""
        return [
            layout.consolidated.parent
            for layout in paths.formats
            if layout.enabled and layout.consolidated is not None
        ]

    def generate(
        self, report_data: Dict[str, Any], paths: ResolvedReportPaths
    ) -> Dict[str, Path]:
        """Write the enabled preflight reports and return them by format."""
        started = time.perf_counter()
        started_at = datetime.now(timezone.utc)

        rows = self._rows(
            report_data,
            PREFLIGHT_VALIDATION_TYPE,
            PREFLIGHT_COLUMNS,
        )

        writers: Dict[str, Callable[..., None]] = {
            "html": self._write_html,
            "json": self._write_json,
            "excel": self._write_excel,
        }

        paths.create_directories()
        generated: Dict[str, Path] = {}

        for layout in paths.formats:
            destination = layout.consolidated

            if not layout.enabled or destination is None:
                logger.info(
                    "run_id=%s phase=%s report_format=%s generation_status=DISABLED",
                    report_data.get("validation_run_id"),
                    paths.phase,
                    layout.name.upper(),
                )
                continue

            # FIX: resolve once and use the exact same filesystem path
            # for directory creation, report writing, and verification.
            output_path = filesystem_path(destination)
            output_path.parent.mkdir(parents=True, exist_ok=True)

            format_started = time.perf_counter()

            try:
                writer = writers.get(layout.name)

                if writer is None:
                    raise ReportPathError(
                        f"Unsupported preflight report format: {layout.name}"
                    )

                writer(
                    output_path,
                    report_data,
                    rows,
                    PREFLIGHT_VALIDATION_TYPE,
                    self.title,
                    PREFLIGHT_COLUMNS,
                    started_at,
                )

            except Exception as exc:
                logger.exception(
                    "run_id=%s phase=%s report_format=%s report_path=%s "
                    "generation_status=FAILED failure_reason=%s duration=%.6f",
                    report_data.get("validation_run_id"),
                    paths.phase,
                    layout.name.upper(),
                    output_path,
                    exc,
                    time.perf_counter() - format_started,
                )
                raise

            if not output_path.is_file():
                raise ReportPathError(
                    f"{layout.name.upper()} preflight report was not written "
                    f"to {output_path}"
                )

            generated[layout.name] = output_path

            logger.info(
                "run_id=%s phase=%s report_format=%s report_path=%s "
                "generation_status=SUCCESS duration=%.6f",
                report_data.get("validation_run_id"),
                paths.phase,
                layout.name.upper(),
                output_path,
                time.perf_counter() - format_started,
            )

        if not generated:
            raise ReportPathError(
                "Every preflight report format is disabled; enable at least "
                "one of report_output.generate_excel/json/html"
            )

        logger.info(
            "run_id=%s phase=%s run_root=%s formats=%s "
            "generation_status=SUCCESS duration=%.6f",
            report_data.get("validation_run_id"),
            paths.phase,
            paths.run_root,
            ",".join(sorted(generated)),
            time.perf_counter() - started,
        )

        return generated

    @staticmethod
    def _rows(data, validation_type, columns):
        run_id = data.get("validation_run_id")
        rows = []
        for raw in data.get("connections", []):
            normalized = dict(raw)
            if validation_type == "preflight" and "table_configuration_id" in raw:
                normalized.update({
                    "overall_validation_status": raw.get(
                        "overall_preflight_validation_status", "FAILED"
                    ),
                    "validation_status": raw.get(
                        "overall_preflight_validation_status", "FAILED"
                    ),
                    "failed_validation_checks": raw.get(
                        "failed_preflight_validation_checks", []
                    ),
                    "failure_reason": raw.get("preflight_failure_reason"),
                    "message": raw.get("preflight_message", ""),
                    "validation_timestamp": raw.get("validation_completion_time"),
                })
            if "table_configuration_id" not in raw:
                side = "source" if raw.get("classification") == "SOURCE" else "target"
                normalized.update({
                    f"{side}_connection_reference": raw.get("connection_reference"),
                    f"{side}_connection_type": raw.get("connection_type"),
                    f"{side}_connection_validation_status": raw.get("validation_status"),
                    f"{side}_object_validation_status": raw.get("validation_status"),
                    "overall_validation_status": raw.get("validation_status"),
                })
            normalized["run_id"] = run_id
            rows.append({column: normalized.get(column) for column in columns})
        return rows

    @staticmethod
    def _summary(rows, report_data=None):
        report_data = report_data or {}
        statuses = [str(row.get("overall_validation_status", "")).upper() for row in rows]
        cumulative = float(
            report_data.get("cumulative_worker_duration") or sum(
                float(row.get("validation_duration_seconds") or 0) for row in rows
            )
        )
        wall_clock = float(report_data.get("wall_clock_duration") or 0)
        if wall_clock <= 0:
            starts = []
            completions = []
            for row in rows:
                try:
                    starts.append(datetime.fromisoformat(
                        str(row.get("validation_start_time")).replace("Z", "+00:00")
                    ))
                    completions.append(datetime.fromisoformat(
                        str(row.get("validation_completion_time")).replace("Z", "+00:00")
                    ))
                except (TypeError, ValueError):
                    continue
            wall_clock = (
                (max(completions) - min(starts)).total_seconds()
                if starts and completions else cumulative
            )
        return {
            "total_tables": len(rows),
            "passed_count": sum(value in {"PASS", "PASSED"} for value in statuses),
            "failed_count": sum(value in {"FAIL", "FAILED"} for value in statuses),
            "warning_count": sum(value in {"WARN", "WARNING", "PARTIALLY_PASSED"} for value in statuses),
            "wall_clock_duration": round(wall_clock, 6),
            "cumulative_worker_duration": round(cumulative, 6),
            # Backward-compatible field: it now means actual elapsed duration.
            "execution_duration": round(wall_clock, 6),
        }

    @classmethod
    def _payload(cls, report_data, rows, validation_type, started_at):
        summary = cls._summary(rows, report_data)
        return {
            "execution_metadata": {
                "run_id": report_data.get("validation_run_id"),
                "validation_type": validation_type,
                "start_time": report_data.get("validation_timestamp"),
                "end_time": datetime.now(timezone.utc).isoformat(),
                "duration": summary["execution_duration"],
                "operation_timings": report_data.get("operation_timings") or [],
                **summary,
            },
            # How many connections the run established and how often they were
            # reused is a preflight result, so it is reported with the rows it
            # was measured over.
            "connection_metrics": report_data.get("connection_metrics") or {},
            "validation_results": rows,
        }

    @classmethod
    def _write_json(cls, path, data, rows, validation_type, title, columns, started_at):
        path.write_text(
            json.dumps(cls._payload(data, rows, validation_type, started_at), indent=2, default=str),
            encoding="utf-8",
        )

    @classmethod
    def _write_html(cls, path, data, rows, validation_type, title, columns, started_at):
        summary = cls._summary(rows, data)
        document = render_preflight_html(
            title=title,
            run_id=data.get("validation_run_id"),
            timestamp=data.get("validation_timestamp"),
            validation_type=validation_type,
            overall_status=data.get("overall_preflight_status"),
            summary=summary,
            connection_metrics=data.get("connection_metrics") or {},
            columns=columns,
            rows=rows,
        )
        path.write_text(document, encoding="utf-8")

    @classmethod
    def _write_excel(cls, path, data, rows, validation_type, title, columns, started_at):
        summary_values = cls._summary(rows, data)
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Summary"
        summary.append(["Metric", "Value"])
        for key, value in {
            "Validation type": validation_type,
            "Run ID": data.get("validation_run_id"),
            **{key.replace("_", " ").title(): value for key, value in summary_values.items()},
            **{
                key.replace("_", " ").title(): (
                    json.dumps(value, default=str)
                    if isinstance(value, (dict, list)) else value
                )
                for key, value in (data.get("connection_metrics") or {}).items()
            },
        }.items():
            summary.append([key, value])
        detail = workbook.create_sheet("Validation_Result")
        detail.append(list(columns))
        for row in rows:
            values = []
            for column in columns:
                value = row[column]
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, default=str)
                elif column in {
                    "validation_start_time", "validation_completion_time"
                } and isinstance(value, str):
                    try:
                        value = datetime.fromisoformat(value.replace("Z", "+00:00"))
                        if value.tzinfo is not None:
                            value = value.replace(tzinfo=None)
                    except ValueError:
                        pass
                values.append(value)
            detail.append(values)
        for column_name in (
            "validation_start_time", "validation_completion_time"
        ):
            if column_name in columns:
                index = columns.index(column_name) + 1
                for row_index in range(2, detail.max_row + 1):
                    detail.cell(row_index, index).number_format = "yyyy-mm-dd hh:mm:ss"
        detail.freeze_panes = "A2"
        detail.auto_filter.ref = detail.dimensions
        for index, column in enumerate(columns, 1):
            width = max(
                [len(str(column))]
                + [len(str(row[column] or "")) for row in rows]
            ) + 2
            detail.column_dimensions[get_column_letter(index)].width = min(width, 60)
        metadata = workbook.create_sheet("Execution_Metadata")
        metadata.append(["Field", "Value"])
        for key, value in {
            "Framework version": "migration-validator_v6",
            "Run ID": data.get("validation_run_id"),
            "Execution environment": "configured-runtime",
            "Report generation time": datetime.now(timezone.utc).isoformat(),
            "Report status": "SUCCESS",
        }.items():
            metadata.append([key, value])
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        workbook.close()
